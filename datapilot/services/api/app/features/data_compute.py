"""Bounded, deterministic data transformations, aggregates and portable exports.

This module never evaluates data as code, SQL, spreadsheet formulas or templates.
Business roles are supplied by the user; column names have no implicit meaning.
"""

import csv
import io
import json
import math
import re
import textwrap
import zipfile
from collections import Counter
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
from itertools import chain
from typing import Any

import openpyxl

ROLES = {"revenue", "cost", "quantity", "date", "category"}
MEASURES = ("revenue", "cost", "quantity")
NUMERIC = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")


def json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return decimal_text(value)
    if isinstance(value, dict):
        return {key: json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(item) for item in value]
    return value


def decimal_value(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    if len(text) > 400 or not NUMERIC.fullmatch(text):
        return None
    try:
        number = Decimal(text)
        if not number.is_finite() or abs(number.adjusted()) > 350:
            return None
        return number
    except (InvalidOperation, ValueError):
        return None


def decimal_text(value: Decimal) -> str:
    if not value:
        return "0"
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def ratio_text(numerator: Decimal, denominator: Decimal) -> str | None:
    if not denominator:
        return None
    with localcontext() as context:
        context.prec = 28
        return decimal_text(numerator / denominator)


def clean_rows(
    rows: list[dict], columns: list[str], operations: list[dict]
) -> tuple[list[dict], dict]:
    """Apply explicit operations to a copy; unknown or lossy operations fail closed."""
    if not 1 <= len(operations) <= 20:
        raise ValueError("Choose between 1 and 20 cleaning operations")
    result = [dict(row) for row in rows]
    names = list(columns)
    changed_cells = 0
    operation_summaries = []
    for operation in operations:
        if not isinstance(operation, dict) or set(operation) - {"op", "column", "value"}:
            raise ValueError("Invalid cleaning operation fields")
        kind, column = operation.get("op"), operation.get("column")
        allowed = {
            "trim",
            "empty_to_null",
            "drop_duplicates",
            "drop_missing",
            "fill_missing",
            "rename",
            "cast_number",
            "lowercase",
            "uppercase",
        }
        if kind not in allowed:
            raise ValueError("Unsupported cleaning operation")
        if column is not None and column not in names:
            raise ValueError("Cleaning column does not exist")
        if kind in {"drop_missing", "fill_missing", "rename", "cast_number"} and not column:
            raise ValueError("This operation requires a column")
        if kind == "drop_duplicates" and column is not None:
            raise ValueError("Duplicate removal compares complete rows")
        before, changes = len(result), 0
        selected = [column] if column is not None else list(names)
        if kind == "drop_duplicates":
            seen = set()
            unique = []
            for row in result:
                key = json.dumps(json_value([row.get(name) for name in names]), sort_keys=True)
                if key not in seen:
                    unique.append(row)
                    seen.add(key)
            result = unique
        elif kind == "drop_missing":
            result = [row for row in result if row.get(column) is not None]
        elif kind == "rename":
            assert isinstance(column, str)
            target = operation.get("value")
            if (
                not isinstance(target, str)
                or not target.strip()
                or len(target) > 120
                or any(ord(c) < 32 or ord(c) == 127 for c in target)
                or (target in names and target != column)
            ):
                raise ValueError("Choose a unique nonblank column name of at most 120 characters")
            if target != column:
                names[names.index(column)] = target
                for row in result:
                    row[target] = row.pop(column)
                changes = len(result)
        else:
            fill = operation.get("value")
            if kind == "fill_missing" and (
                "value" not in operation
                or fill is None
                or isinstance(fill, (dict, list))
                or (isinstance(fill, str) and len(fill) > 32767)
                or (isinstance(fill, float) and not math.isfinite(fill))
            ):
                raise ValueError("Provide a finite scalar replacement value")
            for row in result:
                for name in selected:
                    old = row.get(name)
                    new = old
                    if kind == "trim" and isinstance(old, str):
                        new = old.strip()
                    elif kind == "empty_to_null" and isinstance(old, str) and not old.strip():
                        new = None
                    elif kind == "lowercase" and isinstance(old, str):
                        new = old.lower()
                    elif kind == "uppercase" and isinstance(old, str):
                        new = old.upper()
                    elif kind == "fill_missing" and old is None:
                        new = fill
                    elif kind == "cast_number" and old is not None:
                        number = decimal_value(old)
                        if number is None:
                            raise ValueError("Numeric conversion would lose a nonnumeric value")
                        if number == number.to_integral_value():
                            new = int(number)
                            if not -(2**63) <= new < 2**63:
                                raise ValueError("Numeric conversion exceeds the integer range")
                        else:
                            new = float(number)
                            if not math.isfinite(new) or Decimal(str(new)) != number:
                                raise ValueError("Numeric conversion would lose decimal precision")
                    if old != new or type(old) is not type(new):
                        row[name] = new
                        changes += 1
        changed_cells += changes
        operation_summaries.append(
            {"op": kind, "changed_cells": changes, "removed_rows": before - len(result)}
        )
    return result, {
        "columns": names,
        "before_rows": len(rows),
        "after_rows": len(result),
        "removed_rows": len(rows) - len(result),
        "changed_cells": changed_cells,
        "operations": operation_summaries,
    }


def iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return None
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        return None


def analyze_rows(rows: list[dict], mapping: dict, filters: dict | None = None) -> dict:
    """Aggregate complete canonical data with decimal arithmetic and coverage evidence."""
    with localcontext() as context:
        # Bounded exponents keep mixed tiny/huge totals exact at this precision.
        context.prec = 1200
        return _analyze_rows(rows, mapping, filters or {})


def validate_analysis(rows: list[dict], mapping: dict, filters: dict) -> None:
    if set(mapping) - ROLES or not mapping:
        raise ValueError("Map at least one supported semantic role")
    if any(not isinstance(column, str) or not column for column in mapping.values()):
        raise ValueError("Mapped roles require column names")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Each semantic role must use a distinct column")
    if rows and any(column not in rows[0] for column in mapping.values()):
        raise ValueError("Mapped column does not exist")
    if set(filters) - {"date_from", "date_to", "categories"}:
        raise ValueError("Unsupported analysis filter")
    date_from, date_to = filters.get("date_from"), filters.get("date_to")
    if any(value is not None and iso_date(value) is None for value in (date_from, date_to)):
        raise ValueError("Date filters must use valid YYYY-MM-DD dates")
    if date_from and date_to and date_from > date_to:
        raise ValueError("The start date must precede the end date")
    if (date_from or date_to) and "date" not in mapping:
        raise ValueError("Date filters require a date mapping")
    categories = filters.get("categories")
    if categories is not None and (
        not isinstance(categories, list)
        or len(categories) > 100
        or any(not isinstance(value, str) or len(value) > 32767 for value in categories)
    ):
        raise ValueError("Categories must be a list of at most 100 text values")
    if categories is not None and "category" not in mapping:
        raise ValueError("Category filters require a category mapping")


def filter_rows(rows: list[dict], mapping: dict, filters: dict | None = None) -> list[dict]:
    filters = filters or {}
    validate_analysis(rows, mapping, filters)
    date_from, date_to = filters.get("date_from"), filters.get("date_to")
    categories = filters.get("categories")
    result = []
    for row in rows:
        day = iso_date(row.get(mapping["date"])) if "date" in mapping else None
        category_value = row.get(mapping["category"]) if "category" in mapping else None
        if date_from and (day is None or day < date_from):
            continue
        if date_to and (day is None or day > date_to):
            continue
        if categories is not None and (
            category_value is None or str(category_value) not in categories
        ):
            continue
        result.append(row)
    return result


def _analyze_rows(rows: list[dict], mapping: dict, filters: dict) -> dict:
    filtered = filter_rows(rows, mapping, filters)
    totals = {role: Decimal(0) for role in MEASURES if role in mapping}
    missing = Counter({role: 0 for role in mapping})
    invalid = Counter({role: 0 for role in mapping})
    present = Counter({role: 0 for role in mapping})
    included, profit_rows, profit = 0, 0, Decimal(0)
    paired_revenue = Decimal(0)
    category_groups: dict[str, dict] = {}
    date_groups: dict[str, dict] = {}

    def group_add(groups, label, values):
        if label not in groups:
            groups[label] = {"label": label, "rows": 0, **{role: Decimal(0) for role in totals}}
            if "revenue" in totals and "cost" in totals:
                groups[label]["profit"] = Decimal(0)
        group = groups[label]
        group["rows"] += 1
        for role, value in values.items():
            if value is not None:
                group[role] += value
        if values.get("revenue") is not None and values.get("cost") is not None:
            group["profit"] += values["revenue"] - values["cost"]

    for row in filtered:
        day = iso_date(row.get(mapping["date"])) if "date" in mapping else None
        category_value = row.get(mapping["category"]) if "category" in mapping else None
        included += 1
        values: dict[str, Decimal | None] = {}
        for role, column in mapping.items():
            raw = row.get(column)
            value: Any
            if raw is None or raw == "":
                missing[role] += 1
                value = None
            elif role in MEASURES:
                value = decimal_value(raw)
                if value is None:
                    invalid[role] += 1
                else:
                    present[role] += 1
                    totals[role] += value
            elif role == "date":
                value = day
                if day is None:
                    invalid[role] += 1
                else:
                    present[role] += 1
            else:
                value = str(raw)
                present[role] += 1
            if role in MEASURES:
                values[role] = value
        revenue_value, cost_value = values.get("revenue"), values.get("cost")
        if revenue_value is not None and cost_value is not None:
            profit += revenue_value - cost_value
            paired_revenue += revenue_value
            profit_rows += 1
        if "category" in mapping:
            group_add(
                category_groups,
                "(missing)" if category_value is None else str(category_value),
                values,
            )
        if day is not None:
            group_add(date_groups, day, values)
    metrics: dict[str, Any] = {
        "rows": included,
        **{role: decimal_text(value) for role, value in totals.items()},
    }
    if "revenue" in totals:
        metrics["average_revenue_per_row"] = ratio_text(
            totals["revenue"], Decimal(present["revenue"])
        )
    if "revenue" in totals and "cost" in totals:
        metrics["profit"] = decimal_text(profit)
        metrics["margin_percent"] = ratio_text(profit * 100, paired_revenue)
    evidence = {
        "rows_total": len(rows),
        "rows_included": included,
        "mapping": mapping,
        "filters": filters,
        "missing_values": dict(missing),
        "invalid_values": dict(invalid),
        "valid_values": dict(present),
        "profit_rows": profit_rows,
        "series_limit": 50,
        "category_groups_total": len(category_groups),
        "date_groups_total": len(date_groups),
        "ratio_precision": 28,
        "calculation_policy": (
            "Totals exclude missing/invalid values. Profit and margin use rows with both "
            "valid revenue and cost. Average revenue uses valid revenue rows. "
            "No currency conversion or inferred joins."
        ),
    }
    insights = [
        {
            "code": "row_coverage",
            "text": f"Analysis includes {included} of {len(rows)} source rows.",
            "evidence": {"rows_included": included, "rows_total": len(rows)},
        }
    ]
    if any(missing.values()) or any(invalid.values()):
        insights.append(
            {
                "code": "incomplete_values",
                "text": (
                    "Some mapped values are missing or invalid; "
                    "review coverage before interpreting totals."
                ),
                "evidence": {"missing_values": dict(missing), "invalid_values": dict(invalid)},
            }
        )
    category_series = sorted(
        category_groups.values(), key=lambda group: (-group["rows"], group["label"])
    )[:50]
    date_series = [date_groups[label] for label in sorted(date_groups)[:50]]
    if category_series:
        insights.append(
            {
                "code": "largest_category",
                "text": "The largest category by row count is shown first.",
                "value": category_series[0]["label"],
                "evidence": {"rows": category_series[0]["rows"], "basis": "row_count"},
            }
        )
    return json_value(
        {
            "mapping": mapping,
            "metrics": metrics,
            "category_series": category_series,
            "date_series": date_series,
            "insights": insights,
            "evidence": evidence,
        }
    )


def spreadsheet_text(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return json_value(value)


def csv_bytes(rows: list[dict], columns: list[str]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([spreadsheet_text(name) for name in columns])
    writer.writerows([[spreadsheet_text(row.get(name)) for name in columns] for row in rows])
    return output.getvalue().encode("utf-8-sig")


def pdf_bytes(analysis: dict | None, row_count: int, columns: list[str]) -> bytes:
    """A fixed English numerical report, using standard PDF fonts without dependencies."""
    analysis = analysis or {}
    lines = [
        "DataPilot analysis report",
        "",
        f"Dataset rows: {row_count}",
        f"Columns: {len(columns)}",
    ]
    for key in ("id", "dataset_version_id", "version_id", "source_sha256", "created_at"):
        if analysis.get(key):
            lines.append(f"{key.replace('_', ' ').title()}: {analysis[key]}")
    lines += ["", "Deterministic metrics"]
    lines += [
        f"{key.replace('_', ' ').title()}: {value if value is not None else 'Unavailable'}"
        for key, value in analysis.get("metrics", {}).items()
    ]
    evidence = analysis.get("evidence", {})
    lines += [
        "",
        "Evidence and calculation policy",
        f"Rows included: {evidence.get('rows_included', row_count)}",
        f"Source rows: {evidence.get('rows_total', row_count)}",
        f"Profit paired rows: {evidence.get('profit_rows', 'Unavailable')}",
    ]
    lines += textwrap.wrap(
        evidence.get("calculation_policy", "No semantic mapping was applied."), width=88
    )
    lines += [
        "",
        "The data exports retain all Unicode source values.",
        "CSV/XLSX formula-like text is escaped for spreadsheet safety.",
    ]
    # User-controlled names/questions are deliberately not a PDF template. This fixed
    # report contains numeric aggregates and version provenance; JSON has full labels.
    objects: list[bytes] = []
    objects.extend(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
    )
    page_refs = []
    for index in range(0, len(lines), 48):
        page_id = len(objects) + 1
        stream_id = page_id + 1
        page_refs.append(f"{page_id} 0 R")
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {stream_id} 0 R >>"
            ).encode()
        )
        stream_lines = ["BT /F1 11 Tf 48 790 Td 15 TL"]
        for line in lines[index : index + 48]:
            safe = (
                line.encode("ascii", "backslashreplace")
                .decode()
                .replace("\\", "\\\\")
                .replace("(", "\\(")
                .replace(")", "\\)")
            )
            stream_lines.append(f"({safe}) Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("ascii")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[1] = (
        f"<< /Type /Pages /Count {len(page_refs)} /Kids [{' '.join(page_refs)}] >>".encode()
    )
    output = io.BytesIO()
    output.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(output.tell())
        output.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = output.tell()
    output.write(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return output.getvalue()


def generate_export(
    format: str, rows: list[dict], columns: list[str], analysis: dict | None = None
) -> tuple[bytes, str, str]:
    """Return a download generated from the selected complete immutable version."""
    if format == "csv":
        return csv_bytes(rows, columns), "text/csv; charset=utf-8", "csv"
    if format == "json":
        payload = {"columns": columns, "rows": rows, "analysis": analysis}
        return (
            json.dumps(json_value(payload), ensure_ascii=False, allow_nan=False, indent=2).encode(),
            "application/json",
            "json",
        )
    if format == "xlsx":
        # Validate before starting openpyxl's write-only XML generator so a rejected
        # dataset leaves no partially open worksheet stream behind.
        value_rows: Iterable[Iterable[Any]] = chain(
            [columns], ([row.get(name) for name in columns] for row in rows)
        )
        for row_values in value_rows:
            if any(
                isinstance(value, str) and re.search(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", value)
                for value in row_values
            ):
                raise ValueError(
                    "XLSX cannot represent a control character in this dataset; use CSV or JSON"
                )
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet("Data")

        def cells(values):
            result = []
            for value in values:
                converted = json_value(value)
                # Excel's numeric cells support only 15 significant decimal digits.
                if isinstance(converted, int) and not isinstance(converted, bool):
                    if len(str(abs(converted))) > 15:
                        converted = str(converted)
                cell = openpyxl.cell.WriteOnlyCell(sheet, value=converted)
                if isinstance(converted, str):
                    cell.data_type = "s"
                result.append(cell)
            return result

        sheet.append(cells(columns))
        for row in rows:
            sheet.append(cells([row.get(name) for name in columns]))
        buffer = io.BytesIO()
        book.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    if format == "pdf":
        return pdf_bytes(analysis, len(rows), columns), "application/pdf", "pdf"
    if format in {"powerbi", "powerbi_zip"}:
        buffer = io.BytesIO()
        mapping = {
            role: spreadsheet_text(name)
            for role, name in (analysis or {}).get("mapping", {}).items()
        }
        if any(any(ord(character) < 32 for character in name) for name in mapping.values()):
            raise ValueError(
                "Power BI column references cannot include control characters; "
                "rename the mapped column first or use JSON"
            )
        measures = ["Source Rows = COUNTROWS(Data)"]
        for role in MEASURES:
            if role in mapping:
                name = mapping[role].replace("]", "]]")
                measures.append(f"Total {role.title()} = SUM(Data[{name}])")
        if "revenue" in mapping and "cost" in mapping:
            revenue = mapping["revenue"].replace("]", "]]")
            cost = mapping["cost"].replace("]", "]]")
            paired = f"FILTER(Data, ISNUMBER(Data[{revenue}]) && ISNUMBER(Data[{cost}]))"
            measures += [
                f"Profit = SUMX({paired}, Data[{revenue}] - Data[{cost}])",
                f"Margin % = DIVIDE([Profit], SUMX({paired}, Data[{revenue}]))",
            ]
        query = (
            'let\n    Source = Csv.Document(File.Contents("C:\\DataPilot\\data.csv"),'
            '[Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),\n'
            "    Headers = Table.PromoteHeaders(Source, [PromoteAllScalars=true])\n"
            "in\n    Headers\n"
        )
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("data.csv", csv_bytes(rows, columns))
            archive.writestr(
                "analysis.json",
                json.dumps(json_value(analysis or {}), ensure_ascii=False, indent=2),
            )
            archive.writestr("mapping.json", json.dumps(mapping, ensure_ascii=False, indent=2))
            archive.writestr("measures.dax", "\n".join(measures) + "\n")
            archive.writestr("PowerQuery.m", query)
            archive.writestr(
                "README.txt",
                "DataPilot Power BI handoff\n\n"
                "1. Extract this ZIP. In Power BI Desktop choose Get Data > Text/CSV > data.csv.\n"
                "2. Promote headers, name the table Data, and set numeric/date column types "
                "after reviewing mapping.json. Replace invalid numeric values with null.\n"
                "3. PowerQuery.m is an optional starter: replace its local example path.\n"
                "4. Add each measures.dax formula as a separate measure. "
                "Review missing-value coverage in analysis.json first.\n"
                "5. Create visuals using the confirmed role mapping and saved analysis metrics.\n\n"
                "This is a portable handoff package, not a PBIX/PBIT file. "
                "CSV text beginning with formula characters is prefixed with an apostrophe "
                "for spreadsheet safety. JSON retains original text. Profit/margin use rows "
                "with both valid revenue and cost. "
                "No refresh credentials or API keys are included.\n",
            )
        return buffer.getvalue(), "application/zip", "zip"
    raise ValueError("Choose csv, xlsx, json, pdf or powerbi")
