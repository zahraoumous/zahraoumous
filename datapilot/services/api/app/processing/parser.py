"""Deterministic CSV/XLSX parser. Called only by the restricted child runner."""

import csv
import hashlib
import io
import math
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

MAX_ROWS = 250_000
MAX_COLUMNS = 50
MAX_BYTES = 25 * 1024**2
MAX_EXPANDED_BYTES = 200 * 1024**2


class ParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def normalized_headers(values) -> list[str]:
    headers = []
    seen: set[str] = set()
    for index, value in enumerate(values):
        name = str(value).strip() if value is not None else ""
        name = name[:120] or f"column_{index + 1}"
        candidate = name
        suffix = 2
        while candidate in seen:
            candidate = f"{name}_{suffix}"
            suffix += 1
        seen.add(candidate)
        headers.append(candidate)
    if not headers:
        raise ParseError("empty_file", "The file does not contain a header row.")
    if len(headers) > MAX_COLUMNS:
        raise ParseError("column_limit", "Files may contain at most 50 columns.")
    return headers


def string_value(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    result = str(value)
    if len(result) > 32767:
        raise ParseError("cell_limit", "A cell exceeds the supported text length.")
    return result if result.strip() else None


def csv_rows(path: Path):
    raw = path.read_bytes()
    if b"\0" in raw:
        raise ParseError("invalid_csv", "Binary or NUL-containing input is not a supported CSV.")
    try:
        decoded = raw.decode("utf-8-sig")
        encoding = "utf-8"
    except UnicodeDecodeError:
        try:
            decoded = raw.decode("cp1252")
            encoding = "windows-1252"
        except UnicodeDecodeError as exc:
            raise ParseError("unsupported_encoding", "Use UTF-8 or Windows-1252 CSV text.") from exc
    if not decoded.strip():
        raise ParseError("empty_file", "The CSV is empty.")
    csv.field_size_limit(32767)
    sample = decoded[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter, strict=True)
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as exc:
        raise ParseError("invalid_csv", "The CSV header could not be parsed.") from exc
    return header, reader, {"encoding": encoding, "delimiter": delimiter}


def validate_xlsx(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise ParseError("zip_bomb", "The spreadsheet expands beyond the supported limits.")
            for entry in entries:
                member = PurePosixPath(entry.filename)
                if member.is_absolute() or ".." in member.parts or "\\" in entry.filename:
                    raise ParseError("invalid_xlsx", "Invalid spreadsheet archive member path.")
                if entry.flag_bits & 1:
                    raise ParseError("invalid_xlsx", "Encrypted spreadsheets are not supported.")
                if entry.file_size > 10 * 1024**2 and entry.file_size > 200 * max(
                    entry.compress_size, 1
                ):
                    raise ParseError("zip_bomb", "The spreadsheet compression ratio is excessive.")
                if "vbaproject" in entry.filename.lower():
                    raise ParseError(
                        "macros_unsupported", "Macro-enabled workbooks are not supported."
                    )
            if "[Content_Types].xml" not in archive.namelist():
                raise ParseError("invalid_xlsx", "The file is not a valid XLSX workbook.")
    except zipfile.BadZipFile as exc:
        raise ParseError("invalid_xlsx", "The file is not a valid XLSX workbook.") from exc


def xlsx_rows(path: Path):
    validate_xlsx(path)
    try:
        # Staged originals deliberately have no client-controlled extension.
        # A bounded binary stream lets openpyxl inspect ZIP content directly.
        workbook = openpyxl.load_workbook(
            io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False
        )
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not visible:
            raise ParseError("invalid_xlsx", "The workbook has no visible worksheets.")
        sheet = visible[0]
        # Do not trust an oversized or deliberately underspecified dimension attribute.
        sheet.reset_dimensions()
        cells = sheet.iter_rows()
        first = next(cells, None)
        if first is None:
            workbook.close()
            raise ParseError("empty_file", "The first visible worksheet is empty.")

        def values(row):
            result = []
            for cell in row:
                if cell.data_type == "f":
                    raise ParseError(
                        "formulas_unsupported", "Save formula cells as values before importing."
                    )
                result.append(cell.value)
            return result

        header = values(first)

        def remaining():
            try:
                for row in cells:
                    yield values(row)
            finally:
                workbook.close()

        return (
            header,
            remaining(),
            {
                "worksheet": sheet.title,
                "sheet_names": workbook.sheetnames,
                "worksheet_policy": "First visible sheet. Import other sheets as separate files.",
            },
        )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("invalid_xlsx", "The spreadsheet could not be read safely.") from exc


def number(value: str) -> Decimal | None:
    try:
        if len(value.strip()) > 100 or not re.fullmatch(
            r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value.strip()
        ):
            return None
        result = Decimal(value.strip())
        return result if result.is_finite() and math.isfinite(float(result)) else None
    except (InvalidOperation, OverflowError, ValueError):
        return None


def infer(values: list[str | None]):
    present = [value for value in values if value is not None]
    if not present:
        return "string", values, {}
    # Preserve identifiers such as 00123 instead of silently losing leading zeros.
    if any(re.fullmatch(r"[+-]?0\d+", value.strip()) for value in present):
        return "string", values, {}
    numeric = [number(value) for value in present]
    if all(value is not None for value in numeric):
        integers = all(re.fullmatch(r"[+-]?\d+", value.strip()) for value in present)
        if integers and all(-(2**63) <= int(value) < 2**63 for value in present):
            integer_values = [None if value is None else int(value) for value in values]
            integer_total = sum(int(value) for value in present)
            return (
                "integer",
                integer_values,
                {
                    "sum": str(integer_total),
                    "min": min(
                        converted_value
                        for converted_value in integer_values
                        if converted_value is not None
                    ),
                    "max": max(
                        converted_value
                        for converted_value in integer_values
                        if converted_value is not None
                    ),
                },
            )
        if integers:
            return "string", values, {}
        converted = [None if value is None else float(value) for value in values]
        total = sum((value for value in numeric if value is not None), Decimal(0))
        return (
            "number",
            converted,
            {
                "sum": str(total),
                "min": min(float(value) for value in present),
                "max": max(float(value) for value in present),
            },
        )
    if all(value.lower() in ("true", "false") for value in present):
        return (
            "boolean",
            [None if value is None else value.lower() == "true" for value in values],
            {},
        )
    try:
        if all(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) for value in present):
            [date.fromisoformat(value) for value in present]
            return (
                "date",
                [None if value is None else date.fromisoformat(value) for value in values],
                {},
            )
    except ValueError:
        pass
    return "string", values, {"numeric_like_count": sum(value is not None for value in numeric)}


def parse(path: Path, kind: str, output: Path) -> dict:
    if not 0 < path.stat().st_size <= MAX_BYTES:
        raise ParseError("file_size", "Files must be nonempty and no larger than 25 MiB.")
    if kind == "csv":
        raw_headers, input_rows, source_info = csv_rows(path)
    elif kind == "xlsx":
        raw_headers, input_rows, source_info = xlsx_rows(path)
    else:
        raise ParseError("unsupported_type", "Only CSV and XLSX files are supported.")
    headers = normalized_headers(raw_headers)
    rows: list[list[str | None]] = []
    try:
        for input_row in input_rows:
            if len(input_row) > len(headers):
                raise ParseError("row_width", "A row contains more columns than the header.")
            row = [string_value(value) for value in input_row]
            if not any(value is not None for value in row):
                continue
            row += [None] * (len(headers) - len(row))
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ParseError(
                    "row_limit", "Files may contain at most 250,000 nonblank data rows."
                )
    except csv.Error as exc:
        raise ParseError(
            "invalid_csv", "The CSV has malformed quoting or an excessive cell size."
        ) from exc
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(
            "invalid_xlsx" if kind == "xlsx" else "invalid_csv", "The file could not be parsed."
        ) from exc
    columns, arrays, quality, statistics = [], {}, [], {}
    formula_rows: set[int] = set()
    for index, name in enumerate(headers):
        values = [row[index] for row in rows]
        typ, converted, profile = infer(values)
        nulls = sum(value is None for value in values)
        columns.append({"name": name, "type": typ, "null_count": nulls})
        statistics[name] = profile
        arrow_type = {
            "integer": pa.int64(),
            "number": pa.float64(),
            "boolean": pa.bool_(),
            "date": pa.date32(),
        }.get(typ, pa.string())
        arrays[name] = pa.array(converted, type=arrow_type)
        if nulls:
            quality.append(
                {
                    "rule_code": "missing_values",
                    "severity": "warning",
                    "column_name": name,
                    "affected_rows": nulls,
                    "details_json": {"fraction": nulls / len(rows)},
                }
            )
        numeric_like = profile.get("numeric_like_count", 0)
        nonnull = len(values) - nulls
        if nonnull and 0.8 <= numeric_like / nonnull < 1:
            quality.append(
                {
                    "rule_code": "mixed_types",
                    "severity": "warning",
                    "column_name": name,
                    "affected_rows": nonnull - numeric_like,
                    "details_json": {
                        "description": "Mostly numeric column includes text; preserved as text."
                    },
                }
            )
        for row_index, value in enumerate(values):
            if value and value.lstrip().startswith(("=", "+", "-", "@")) and number(value) is None:
                formula_rows.add(row_index)
    duplicate_rows = sum(count - 1 for count in Counter(tuple(row) for row in rows).values())
    if duplicate_rows:
        quality.append(
            {
                "rule_code": "duplicate_rows",
                "severity": "warning",
                "column_name": None,
                "affected_rows": duplicate_rows,
                "details_json": {
                    "description": "Exact duplicate occurrences beyond the first; no rows removed."
                },
            }
        )
    if formula_rows:
        quality.append(
            {
                "rule_code": "formula_like_values",
                "severity": "info",
                "column_name": None,
                "affected_rows": len(formula_rows),
                "details_json": {
                    "description": "Formula-like strings were kept as text, never evaluated."
                },
            }
        )
    table = pa.table(arrays)
    pq.write_table(table, output, compression="snappy")
    preview = table.slice(0, 100).to_pylist()
    for row in preview:
        for key, value in row.items():
            if isinstance(value, (date, datetime)):
                row[key] = value.isoformat()
            elif isinstance(value, str) and len(value) > 256:
                row[key] = value[:256] + "…"
    return {
        "row_count": len(rows),
        "column_count": len(headers),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "schema_json": {
            "columns": columns,
            "preview": preview,
            "profile": {
                "duplicate_rows": duplicate_rows,
                "statistics": statistics,
                **source_info,
                "preview_string_limit": 256,
                "rows_skipped_policy": "Entirely blank rows are skipped; original bytes preserved.",
            },
        },
        "quality": quality,
    }
